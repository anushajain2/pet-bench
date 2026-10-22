"""Audited cross-study ranking. External outcomes never enter model fitting."""
import argparse
import hashlib
import itertools
import json
from pathlib import Path
import numpy as np
import pandas as pd
from Bio.Seq import Seq
from openpyxl import load_workbook
from scipy.stats import spearmanr
from threadpoolctl import threadpool_limits
from .bench_eval import validate_data_inputs
from .campaign import bootstrap_predict
from .data import digest, write_json
from .features import composition, make_aligner, alignment_identity, sequence_fingerprint
from .inference import embed_sequences, MODEL_ID, REVISION


def translate_insert(dna):
    dna = ''.join(str(dna).split()).upper()
    if set(dna) - set('ACGT') or not dna.startswith('CATATG') or not dna.endswith('TAAGCGGCCGC'):
        raise ValueError('Expected documented NdeI / stop / NotI construct boundaries.')
    cds = dna[3:-11]  # terminal TAA + eight-base NotI site
    if len(cds) % 3:
        raise ValueError('Incomplete codon in coding insert.')
    protein = str(Seq(cds).translate())
    if not protein.startswith('M') or '*' in protein:
        raise ValueError('Invalid coding insert translation.')
    return protein


def import_study(root, config):
    raw = root/'external/raw'
    wb = load_workbook(raw/config['sequence_file'], read_only=True, data_only=True)
    sequences = []
    for name, dna in list(wb.active.values)[1:]:
        seq = translate_insert(dna)
        sequences.append({'sample_id': name, 'sequence': seq, 'length': len(seq),
                          'sequence_sha256': hashlib.sha256(seq.encode()).hexdigest()})
    wb.close()
    table = pd.DataFrame(sequences)
    if len(table) != 34 or table.sample_id.duplicated().any() or table.sequence.duplicated().any():
        raise ValueError('Expected 34 distinct APET inserts.')
    wb = load_workbook(raw/config['assay_file'], read_only=True, data_only=True)
    rows = []
    temperature = None
    for row_number, r in enumerate(list(wb[config['sheet']].values)[3:], start=4):
        temperature = r[0] if r[0] is not None else temperature
        if not str(r[1]).startswith('APET-'):
            continue
        values = np.asarray(r[2:5], dtype=float)
        if not np.isfinite(values).all() or (values < 0).any():
            raise ValueError(f'Missing/invalid PET replicate in source row {row_number}.')
        if not np.isclose(values.mean(), float(r[5]), rtol=1e-5, atol=1e-6):
            raise ValueError('Published mean does not match replicates.')
        rows.append({'sample_id': r[1], 'temperature_c': int(temperature),
                     'activity_um': values.mean(), 'sd_um': values.std(ddof=1),
                     'replicate_1': values[0], 'replicate_2': values[1], 'replicate_3': values[2],
                     'source_row': row_number})
    wb.close()
    measured = pd.DataFrame(rows)
    if measured.duplicated(['sample_id', 'temperature_c']).any() or len(measured) != 70:
        raise ValueError('Expected 14 candidates at five temperatures.')
    return table, measured.merge(table, on='sample_id', validate='many_to_one')


def construct_overlap(a, b):
    return min(len(a), len(b))/max(len(a), len(b)) >= .9 and (a in b or b in a)


def cross_identity(external, training):
    aligner = make_aligner()
    return np.array([[alignment_identity(a, b, aligner) for b in training.sequence]
                     for a in external.sequence])


def evaluate_scores(ids, y, scores, k):
    order = np.lexsort((np.asarray(ids, dtype=str), -np.asarray(scores)))
    selected = order[:k]
    random_means = np.asarray([np.mean(y[list(c)]) for c in itertools.combinations(range(len(y)), k)])
    selected_mean = float(y[selected].mean())
    corr = float(spearmanr(y, scores).statistic) if np.ptp(scores) > 0 and np.ptp(y) > 0 else np.nan
    return {'n': len(y), 'top_k': k, 'spearman': corr, 'selected_mean_um': selected_mean,
            'random_mean_um': float(y.mean()), 'random_p05_um': float(np.quantile(random_means, .05)),
            'random_p95_um': float(np.quantile(random_means, .95)),
            'random_fraction_at_least_selected': float(np.mean(random_means >= selected_mean-1e-10)),
            'oracle_mean_um': float(np.sort(y)[-k:].mean()),
            'selected_ids': ', '.join(np.asarray(ids)[selected])}


def run(root):
    config = json.loads((root/'external_protocol.json').read_text())
    validate_data_inputs(root, json.loads((root/'benchmark_config.json').read_text()))
    base = root/'external'
    base.mkdir(exist_ok=True)
    sequences, measured = import_study(root, config)
    catalog = pd.read_csv(root/'benchmark/proteins.csv').sort_values('sample_id').reset_index(drop=True)
    external = sequences[sequences.sample_id.isin(measured.sample_id)].sort_values('sample_id').reset_index(drop=True)
    similarity = cross_identity(external, catalog)
    audit = external[['sample_id', 'length', 'sequence_sha256']].copy()
    nearest = similarity.argmax(axis=1)
    audit['nearest_training_id'] = catalog.sample_id.to_numpy()[nearest]
    audit['max_training_identity'] = similarity.max(axis=1)
    audit['construct_overlap_ids'] = [', '.join(catalog.loc[[construct_overlap(s, t) for t in catalog.sequence], 'sample_id'])
                                      for s in external.sequence]
    audit['excluded_overlap'] = audit.construct_overlap_ids.ne('')
    audit.to_csv(base/'sequence_audit.csv', index=False)
    sequences.to_csv(base/'translated_sequences.csv', index=False)
    measured.to_csv(base/'measurements.csv', index=False)
    # Public external vectors may be cached, with exact sequence and model provenance.
    xp = base/'embeddings.npz'
    mp = base/'embeddings.json'
    expected = {'sequence_sha256': sequence_fingerprint(external), 'model_id': MODEL_ID, 'revision': REVISION}
    if xp.exists() and mp.exists():
        meta = json.loads(mp.read_text())
        if any(meta.get(k) != v for k, v in expected.items()) or digest(xp) != meta['sha256']:
            raise ValueError('External embedding cache mismatch.')
        with np.load(xp, allow_pickle=False) as data:
            ex = data['embeddings']
            if not np.array_equal(data['ids'], external.sample_id):
                raise ValueError('External embedding order mismatch.')
    else:
        ex = embed_sequences(root, external.sequence)
        np.savez_compressed(xp, embeddings=ex, ids=external.sample_id.to_numpy(dtype=str))
        write_json(mp, {**expected, 'sha256': digest(xp)})
    if ex.shape != (len(external), 480) or not np.isfinite(ex).all():
        raise ValueError('Invalid external embeddings.')
    training_vectors = embed_sequences(root, catalog.sequence)
    all_rows, metrics = [], []
    for pair in config['pairs']:
        train = pd.read_csv(root/'benchmark/tasks'/pair['training_task']/'candidates.csv').sort_values('sample_id').reset_index(drop=True)
        idx = catalog.set_index('sample_id').index.get_indexer(train.sample_id)
        if (idx < 0).any():
            raise ValueError('Training ID absent from catalog.')
        n = len(train)
        features = {'ESM-2 35M': np.vstack([training_vectors[idx], ex]),
                    'Composition': composition(pd.concat([train, external], ignore_index=True))}
        scores = {}
        with threadpool_limits(limits=1):
            for method, x in features.items():
                scores[method], _ = bootstrap_predict(x, np.arange(n), train.activity.to_numpy(),
                                                       np.arange(n, n+len(external)), config['seed'], config)
        nearest_task = similarity[:, idx].argmax(axis=1)
        scores['Nearest sequence'] = train.activity.to_numpy()[nearest_task]
        # Only now join external labels for evaluation; model code never receives them.
        outcomes = measured[measured.temperature_c.eq(pair['external_temperature_c'])].set_index('sample_id').loc[external.sample_id]
        y = outcomes.activity_um.to_numpy()
        masks = {'nonoverlapping': ~audit.excluded_overlap.to_numpy()}
        low = masks['nonoverlapping'] & audit.max_training_identity.lt(.6).to_numpy()
        if low.sum() >= 6:
            masks['identity_below_60pct'] = low
        for method, score in scores.items():
            for i, sample in enumerate(external.sample_id):
                all_rows.append({**pair, 'model': method, 'sample_id': sample,
                    'score': float(score[i]), 'activity_um': float(y[i]),
                    'excluded_overlap': bool(audit.excluded_overlap.iloc[i]),
                    'max_training_identity': float(audit.max_training_identity.iloc[i])})
            for cohort, mask in masks.items():
                if mask.sum() < 6:
                    continue
                metrics.append({**pair, 'cohort': cohort, 'model': method,
                    **evaluate_scores(external.sample_id.to_numpy()[mask], y[mask], score[mask], config['top_k'])})
    pd.DataFrame(all_rows).to_csv(base/'predictions.csv', index=False)
    pd.DataFrame(metrics).to_csv(base/'metrics.csv', index=False)
    paths = [root/'external_protocol.json', root/'benchmark/data_manifest.json',
             root/'benchmark/embeddings/esm2_35m.npz', root/'benchmark/embeddings/esm2_35m.json',
             root/'src/pet_scout/external.py', root/'src/pet_scout/inference.py',
             root/'src/pet_scout/campaign.py', root/'src/pet_scout/features.py']
    paths += list((base/'raw').glob('*')) + list(base.glob('*.csv')) + [xp, mp]
    write_json(base/'manifest.json', {'source_url': config['source_archive'],
        'source_doi': config['source_doi'], 'source_license': config['source_license'],
        'input_and_output_sha256': {str(p.relative_to(root)): digest(p) for p in paths},
        'n_measured_candidates': len(external), 'n_excluded_overlap': int(audit.excluded_overlap.sum()),
        'n_below_60pct_identity': int(low.sum()),
        'interpretation': config['restrictions']})
    print(pd.DataFrame(metrics).to_string(index=False))


def verified(root):
    manifest = json.loads((root/'external/manifest.json').read_text())
    for name, expected in manifest['input_and_output_sha256'].items():
        if digest(root/name) != expected:
            raise ValueError(f'External result is stale: {name}. Rerun the external evaluation.')
    validate_data_inputs(root, json.loads((root/'benchmark_config.json').read_text()))
    return manifest


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root', type=Path, default=Path.cwd())
    run(parser.parse_args().root.resolve())
